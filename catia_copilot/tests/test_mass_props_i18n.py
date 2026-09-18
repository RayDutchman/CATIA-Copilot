"""mass_props_dialog 的 i18n 数据流测试（Phase 2 Task 2.3 + 审查修复）。

覆盖核心正确性边界：
1. 列头经 _column_header 工厂：单位后缀串接用 translate().format()，语言缺失时回退中文。
2. _display_headers 全量列头均走工厂，返回恒中文。
3. 导出路径（CSV / Excel）在安装真实临时 en_US translator 下仍恒中文数据：
   表头列名经 _export_column_header 不经翻译直接落盘（含当前单位后缀）、
   Density<0 输出"不统一"、汇总行"总计 (根产品)" —— 英文界面下导出仍为中文。
"""
import os
import subprocess
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

import openpyxl  # noqa: E402
from PySide6.QtCore import QCoreApplication, QTranslator  # noqa: E402
from PySide6.QtWidgets import QApplication  # noqa: E402

from catia_copilot.ui.mass_props_dialog import MassPropsDialog  # noqa: E402

_QAPP = QApplication.instance() or QApplication([])


# 真实临时 en_US qm 需覆盖的关键词条（导出表头/数据恒中文路径的对照翻译）
_EN_QM_ENTRIES = [
    ("零件编号", "Part No."),
    ("类型", "Type"),
    ("重量 ({0})", "Weight ({0})"),
    ("密度 (kg/m³)", "Density (kg/m3)"),
    ("不统一", "Inconsistent"),
    ("总计 (根产品)", "Total (root product)"),
    (" (对称件)", " (Mirror)"),
    ("(虚拟)", "(Virtual)"),
]


def _build_tmp_en_qm() -> Path:
    """用真实 pyside6-lrelease 编译仅含上述词条的临时 en_US qm。"""
    d = Path(tempfile.mkdtemp())
    ts = d / "tmp_en.ts"
    msgs = "".join(
        f'<message><source>{src}</source><translation>{tr}</translation></message>\n'
        for src, tr in _EN_QM_ENTRIES
    )
    ts.write_text(
        '<?xml version="1.0" encoding="utf-8"?>\n'
        '<TS version="2.1" language="en_US">\n'
        '<context><name>CATIACopilot</name>\n'
        f"{msgs}"
        "</context></TS>",
        encoding="utf-8",
    )
    subprocess.run(["pyside6-lrelease", str(ts)], check=True, capture_output=True)
    return d / "tmp_en.qm"


class TestColumnHeaders(unittest.TestCase):
    def test_density_header_is_chinese_literal(self) -> None:
        dlg = MassPropsDialog()
        self.assertEqual(dlg._column_header("Density"), "密度 (kg/m³)")

    def test_weight_header_contains_unit_suffix(self) -> None:
        dlg = MassPropsDialog()
        dlg._mass_unit = "kg"
        self.assertEqual(dlg._column_header("Weight"), "重量 (kg)")

    def test_inertia_header_contains_multiplied_unit(self) -> None:
        dlg = MassPropsDialog()
        dlg._inertia_unit = "kg\u00b7mm\u00b2"
        self.assertEqual(dlg._column_header("Ixx"), "Ixx (kg·mm²)")

    def test_cog_header_contains_unit_suffix(self) -> None:
        dlg = MassPropsDialog()
        dlg._cog_unit = "mm"
        self.assertEqual(dlg._column_header("CogX"), "CogX (mm)")

    def test_plain_column_falls_back_to_display_factory(self) -> None:
        dlg = MassPropsDialog()
        self.assertEqual(dlg._column_header("Part Number"), "零件编号")

    def test_display_headers_all_factory(self) -> None:
        dlg = MassPropsDialog()
        dlg._columns = ["Part Number", "Type", "Weight", "Ixx", "CogX"]
        self.assertEqual(
            dlg._display_headers(),
            ["零件编号", "类型", "重量 (g)", "Ixx (g·mm²)", "CogX (mm)"],
        )


class TestExportKeepsChineseUnderEnTranslator(unittest.TestCase):
    """安装真实临时 en_US translator 后，CSV 与 Excel 导出路径仍必须恒中文。"""

    @classmethod
    def setUpClass(cls) -> None:
        try:
            cls._qm = _build_tmp_en_qm()
        except Exception as exc:  # lrelease 缺失等情况
            raise unittest.SkipTest(f"pyside6-lrelease 不可用：{exc}")

    def setUp(self) -> None:
        self._tr = QTranslator(_QAPP)
        loaded = self._tr.load(str(self._qm))
        if not loaded:
            self.fail("临时 qm 加载失败")
        _QAPP.installTranslator(self._tr)

    def tearDown(self) -> None:
        _QAPP.removeTranslator(self._tr)
        self._tr.deleteLater()

    def test_translator_really_active(self) -> None:
        """前置校验：临时 en translator 确已生效，保证导出恒中文断言有意义。"""
        self.assertEqual(
            QCoreApplication.translate("CATIACopilot", "零件编号"), "Part No."
        )
        self.assertEqual(
            QCoreApplication.translate("CATIACopilot", "重量 ({0})").format("g"),
            "Weight (g)",
        )

    def _make_dlg(self) -> MassPropsDialog:
        dlg = MassPropsDialog()
        dlg._build_hierarchy_columns = lambda: [  # noqa: E731
            "Part Number", "Type", "Density", "Weight",
        ]
        dlg._get_hierarchy_rows = lambda: [  # noqa: E731
            {"Part Number": "P1", "Type": "零件", "Density": -1.0, "Weight": 1.2},
            # 镜像行（对称件）：数据层持中文后缀，导出必须保持中文
            {"Part Number": "M1 (对称件)", "Type": "MIRROR",
             "Density": 1.0, "Weight": 1.0},
        ]
        dlg._unit_factor = 1.0
        dlg._inertia_unit_factor = 1.0
        dlg._cog_unit_factor = 1.0
        dlg._mass_unit = "g"
        dlg._inertia_unit = "g\u00b7mm\u00b2"
        dlg._cog_unit = "mm"
        dlg._rollup_result = {"total_weight": 1.2, "cog": [0.0, 0.0, 0.0],
                              "inertia": [[0.0] * 3 for _ in range(3)]}
        return dlg

    def test_csv_export_stays_chinese(self) -> None:
        dlg = self._make_dlg()
        with patch("catia_copilot.ui.mass_props_dialog.MassPropsDialog._row_status",
                   return_value=""):
            dest = Path(__file__).parent / "_tmp_mp_export.csv"
            dest.write_text("", encoding="utf-8")
            try:
                dlg._do_export_csv(dest)
                content = dest.read_text(encoding="utf-8-sig")
                # 恒中文表头与数据（含度更当前单位后缀）
                self.assertIn("零件编号", content)
                self.assertIn("重量 (g)", content)
                self.assertIn("不统一", content)
                self.assertIn("总计 (根产品)", content)
                # 镜像行 PN 数据层保持中文后缀（_make_item 显示翻译不得污染导出）
                self.assertIn("M1 (对称件)", content)
                # 若导出误走了 _column_header，英文下会落盘英文
                self.assertNotIn("Part No.", content)
                self.assertNotIn("Inconsistent", content)
                self.assertNotIn("Total", content)
            finally:
                dest.unlink(missing_ok=True)

    def test_excel_export_header_stays_chinese(self) -> None:
        dlg = self._make_dlg()
        with patch("catia_copilot.ui.mass_props_dialog.MassPropsDialog._row_status",
                   return_value=""):
            dest = Path(__file__).parent / "_tmp_mp_export.xlsx"
            try:
                dlg._do_export(str(dest))
                wb  = openpyxl.load_workbook(dest)
                ws  = wb.active
                header = [ws.cell(row=1, column=ci).value
                          for ci in range(1, 6)]
                self.assertEqual(header,
                                 ["零件编号", "类型", "密度 (kg/m³)", "重量 (g)", "状态"])
                # 数据行 Density<0 恒 "不统一"；镜像行 PN 数据层保持中文后缀；
                # 汇总行恒 "总计 (根产品)"
                self.assertEqual(ws.cell(row=2, column=3).value, "不统一")
                self.assertEqual(ws.cell(row=3, column=1).value, "M1 (对称件)")
                # Type 列导出沿用既有 TYPE_DISPLAY_NAMES 映射（MIRROR 未收录回退英文 key）
                self.assertEqual(ws.cell(row=3, column=2).value, "MIRROR")
                self.assertEqual(ws.cell(row=4, column=1).value, "总计 (根产品)")
                self.assertNotIn("Part No.", str(header))
            finally:
                dest.unlink(missing_ok=True)


class TestMirrorRowDisplayLayer(unittest.TestCase):
    """镜像行（对称件）：数据层保持中文后缀，仅显示层随语言翻译。"""

    @classmethod
    def setUpClass(cls) -> None:
        try:
            cls._qm = _build_tmp_en_qm()
        except Exception as exc:  # lrelease 缺失等情况
            raise unittest.SkipTest(f"pyside6-lrelease 不可用：{exc}")

    @staticmethod
    def _mirror_row() -> dict:
        return {
            "Part Number":   "P1 (对称件)",
            "Instance Name": "INST-1 (对称件)",
            "Filename":      "(虚拟)",
            "Type":          "MIRROR",
            "_is_mirror":    True,
            "_filepath":     "",
            "_not_found":    False,
            "_no_file":      False,
            "_unreadable":   False,
            "_meas_failed":  False,
        }

    def _make_item(self, dlg: MassPropsDialog, row: dict):
        dlg._columns = ["Part Number", "Instance Name", "Filename", "Type"]
        dlg._pn_to_items = {}
        dlg._item_by_row = []
        return dlg._make_item(0, row)

    def test_zh_fallback_keeps_chinese(self) -> None:
        dlg = MassPropsDialog()
        row = self._mirror_row()
        item = self._make_item(dlg, row)
        # 显示层：无翻译器时回退中文后缀
        self.assertEqual(item.text(0), "P1 (对称件)")
        self.assertEqual(item.text(1), "INST-1 (对称件)")
        self.assertEqual(item.text(2), "(虚拟)")
        # 数据层不被视图渲染改动（导出恒中文依赖）
        self.assertEqual(row["Part Number"], "P1 (对称件)")
        self.assertEqual(row["Filename"], "(虚拟)")

    def test_en_translates_display_only(self) -> None:
        row = self._mirror_row()
        tr = QTranslator(_QAPP)
        try:
            self.assertTrue(tr.load(str(self._qm)), "临时 qm 加载失败")
            _QAPP.installTranslator(tr)
            dlg = MassPropsDialog()
            item = self._make_item(dlg, row)
            # 显示层：英文界面下后缀被翻译
            self.assertEqual(item.text(0), "P1 (Mirror)")
            self.assertEqual(item.text(1), "INST-1 (Mirror)")
            self.assertEqual(item.text(2), "(Virtual)")
            self.assertNotIn("(对称件)", [item.text(0), item.text(1)],
                             "英文界面不得残留中文镜像后缀")
            self.assertNotIn("(虚拟)", [item.text(2)],
                             "英文界面不得残留中文虚拟文件名")
        finally:
            _QAPP.removeTranslator(tr)
            tr.deleteLater()
        # 数据层保持中文后缀（显示翻译不得污染 row_data）
        self.assertEqual(row["Part Number"], "P1 (对称件)")
        self.assertEqual(row["Filename"], "(虚拟)")


if __name__ == "__main__":
    unittest.main()